"""Independent-library verification for `add_footnote`/`edit_footnote`.

Companion to `tests/integration/test_paragraph_edit_python_docx_roundtrip.py`
(see that module's docstring for the full rationale). `add_footnote`'s
from-scratch path is the highest-risk write path in the project for exactly
this kind of blind spot: when a document has no footnotes yet, it builds
`word/footnotes.xml`, a `[Content_Types].xml` `Override`, and a
`word/_rels/document.xml.rels` `Relationship` from nothing, in one atomic
step (see
[ADR-0007](../../docs/adr/0007-phase-6-footnote-edit-module-layout-and-multi-part-atomic-write.md)).
Every existing test for that path re-parses the result with this project's
*own* `docx_mcp.footnotes`/`docx_mcp.structure` code - a mistake shared
between the writer and that reader (a wrong namespace, a missing OPC
relationship attribute) could look consistent to both while producing a
package a real OOXML consumer rejects or silently ignores. `python-docx`
does not expose a high-level footnotes API, but its `Package`/`_Relationship`
object model is a fully independent OPC-graph reader: if the from-scratch
wiring were wrong, opening the document at all, or resolving the
`footnotes` relationship from `word/_rels/document.xml.rels`, would fail.

Only real, full-OPC-package fixtures (`structured.docx`, `footnotes.docx`)
are used here, for the same reason
`test_paragraph_edit_python_docx_roundtrip.py` avoids the minimal synthetic
`word/document.xml`-only archives used elsewhere in `tests/test_footnote_edit.py`
for edge cases: those are not valid OPC packages `python-docx` (or Word) can
open at all.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

from docx_mcp.footnote_edit import add_footnote_to_document, edit_footnote_in_document

WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_NSMAP = {"w": WORD_NS}


def _footnote_text_by_id(footnotes_xml: bytes) -> dict[str, str]:
    root = etree.fromstring(footnotes_xml)
    result: dict[str, str] = {}
    for footnote in root.findall("w:footnote", namespaces=_NSMAP):
        footnote_id = footnote.get(f"{{{WORD_NS}}}id")
        assert footnote_id is not None
        texts = [t.text or "" for t in footnote.findall(".//w:t", namespaces=_NSMAP)]
        result[footnote_id] = "".join(texts)
    return result


def test_from_scratch_footnote_package_reopens_and_resolves_via_independent_opc_graph(
    structured_docx: Path,
) -> None:
    """`structured.docx` has no footnotes yet, so this exercises the
    from-scratch `word/footnotes.xml` + `[Content_Types].xml` +
    `word/_rels/document.xml.rels` creation path."""
    new_id = add_footnote_to_document(structured_docx, 1, "A brand new footnote.")

    document = Document(str(structured_docx))

    footnotes_relationships = [
        rel for rel in document.part.rels.values() if rel.reltype == RT.FOOTNOTES
    ]
    assert len(footnotes_relationships) == 1
    footnotes_part = footnotes_relationships[0].target_part
    assert str(footnotes_part.partname) == "/word/footnotes.xml"
    assert footnotes_part.content_type == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
    )

    content_by_id = _footnote_text_by_id(footnotes_part.blob)
    assert content_by_id[new_id] == "A brand new footnote."


def test_footnote_added_to_a_document_that_already_has_footnotes_reopens_correctly(
    footnotes_docx: Path,
) -> None:
    new_id = add_footnote_to_document(footnotes_docx, 5, "Fourth footnote.")

    document = Document(str(footnotes_docx))
    footnotes_part = next(
        rel.target_part for rel in document.part.rels.values() if rel.reltype == RT.FOOTNOTES
    )

    content_by_id = _footnote_text_by_id(footnotes_part.blob)
    assert content_by_id["1"] == " First footnote."
    assert content_by_id["2"] == " Second footnote."
    assert content_by_id["3"] == " Third footnote."
    assert content_by_id[new_id] == "Fourth footnote."


def test_edited_footnote_reopens_correctly_with_updated_content(footnotes_docx: Path) -> None:
    edit_footnote_in_document(footnotes_docx, "2", "Rewritten second footnote.")

    document = Document(str(footnotes_docx))
    footnotes_part = next(
        rel.target_part for rel in document.part.rels.values() if rel.reltype == RT.FOOTNOTES
    )

    content_by_id = _footnote_text_by_id(footnotes_part.blob)
    assert content_by_id["1"] == " First footnote."
    assert content_by_id["2"] == "Rewritten second footnote."
    assert content_by_id["3"] == " Third footnote."


def test_document_body_text_is_still_readable_after_from_scratch_footnote_creation(
    structured_docx: Path,
) -> None:
    """The from-scratch path also rewrites `word/document.xml` (to append the
    new `w:footnoteReference`) in the same atomic multi-part write - this
    confirms that rewrite doesn't disturb the rest of the body when read back
    through python-docx's own paragraph model, independent of
    `docx_mcp.document.extract_text`."""
    add_footnote_to_document(structured_docx, 1, "A brand new footnote.")

    document = Document(str(structured_docx))

    assert [p.text for p in document.paragraphs] == [
        "Introduction",
        "This is a plain paragraph under the introduction.",
        "Background",
        "Custom styled heading text",
        "Deep Dive",
        "Final plain paragraph.",
    ]
