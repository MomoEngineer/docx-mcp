"""Independent-library verification for `replace_text`.

Companion to `tests/integration/test_paragraph_edit_python_docx_roundtrip.py`
(see that module's docstring for the full rationale): every other test for
`replace_text` verifies its output using this project's *own* parsing code
(`text_edit.py`'s helpers, reused by `find_text`/`get_structure`) - a blind
spot there could produce a document that looks correct to our own assertions
while being subtly wrong to a real OOXML consumer. `replace_text`'s
run-splitting/merging algorithm is the most XML-surgery-heavy read/write path
in the project (it creates new `w:r`/`w:t` siblings, deep-copies `w:rPr`, and
removes emptied runs), which makes it exactly the kind of write path this
independent check is for. Added during the same post-implementation
verification pass as the paragraph-edit round-trip module, for the same
reason: closing a documented blind spot, not responding to a spec change.

Only `run_split.docx` (a real, full-OPC-package `python-docx` fixture) is
used here, for the same reason `test_paragraph_edit_python_docx_roundtrip.py`
only uses `paragraph_edits.docx`: the minimal synthetic `word/document.xml`
-only archives used elsewhere in `tests/test_text_edit.py` for edge cases are
not valid OPC packages `python-docx` (or Word) can open at all.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from docx_mcp.text_edit import TextLocation, replace_text_in_document


def test_edited_document_reopens_correctly_in_python_docx(run_split_docx: Path) -> None:
    replace_text_in_document(run_split_docx, "receive", "REPLIED", location=TextLocation(0, 7, 14))
    replace_text_in_document(run_split_docx, "test", "exam")

    document = Document(str(run_split_docx))

    assert [p.text for p in document.paragraphs] == [
        "Please REPLIED this message.",
        "This is a simple example sentence.",
        "Left\tRight",
        "This paragraph mentions exam, exam, and exam again.",
        "One more exam paragraph for global replace.",
    ]


def test_boundary_run_formatting_survives_reopening_in_python_docx(
    run_split_docx: Path,
) -> None:
    """`run_split.docx`'s first paragraph splits "receive" across a bold run
    (`"rec"`) and a not-bold run (`"eive"`) - independent confirmation, via
    python-docx's own `Run.bold`, that a location-mode replacement inherits
    the *leftmost* matched run's formatting and leaves the untouched prefix
    ("Please ") and suffix (" this message.") runs' formatting alone."""
    replace_text_in_document(run_split_docx, "receive", "REPLIED", location=TextLocation(0, 7, 14))

    document = Document(str(run_split_docx))
    paragraph = document.paragraphs[0]

    runs_by_text = {run.text: run for run in paragraph.runs}
    assert runs_by_text["Please "].bold is None
    assert runs_by_text["REPLIED"].bold is True
    assert runs_by_text[" this message."].bold is None


def test_replacement_inside_italic_run_stays_italic_when_reopened(
    run_split_docx: Path,
) -> None:
    """The second paragraph's whole text is one italic run - independent
    confirmation that the trivial single-run case's replacement run carries
    the original run's `rPr` (here, italic), not default formatting."""
    replace_text_in_document(run_split_docx, "simple", "brief", location=TextLocation(1, 10, 16))

    document = Document(str(run_split_docx))
    paragraph = document.paragraphs[1]

    assert paragraph.text == "This is a brief example sentence."
    assert all(run.italic for run in paragraph.runs if run.text)


def test_untouched_paragraphs_keep_their_formatting_after_an_unrelated_edit(
    run_split_docx: Path,
) -> None:
    """A location-mode edit to paragraph 0 must not disturb paragraph 1's
    independently-observable italic formatting - the formatting-fidelity
    guarantee (docs/testing.md §2.3), checked through a second library."""
    replace_text_in_document(run_split_docx, "receive", "REPLIED", location=TextLocation(0, 7, 14))

    document = Document(str(run_split_docx))
    paragraph = document.paragraphs[1]

    assert paragraph.text == "This is a simple example sentence."
    assert all(run.italic for run in paragraph.runs if run.text)
