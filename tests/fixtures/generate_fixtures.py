"""Generates the synthetic `.docx` fixtures under `tests/fixtures/`.

Dev-only tooling (see [ADR-0001](../../docs/adr/0001-ooxml-library-and-module-layout.md),
Alternatives): uses `python-docx` to build valid base OOXML reliably, then
performs the same kind of raw OOXML surgery `docx-mcp` itself does to splice
in a footnote, since `python-docx` (as of the pinned version) has no
high-level footnote API. Never imported by `src/docx_mcp/`. Re-run with:

    python tests/fixtures/generate_fixtures.py

to regenerate `minimal.docx`/`structured.docx`/`footnotes.docx` after changing
their expected structure (keep the corresponding test file's expectations -
`tests/test_document.py`, `tests/test_structure.py`, `tests/test_footnotes.py`
- in sync manually; this script is not run automatically by the test suite).
"""

from __future__ import annotations

import datetime
import shutil
import zipfile
from collections.abc import Sequence
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree

FIXTURES_DIR = Path(__file__).parent
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
FOOTNOTES_RELATIONSHIP_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
)
FOOTNOTES_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
)
W = f"{{{WORD_NS}}}"


def _w(tag: str) -> str:
    return f"{W}{tag}"


def _build_base_docx(path: Path) -> None:
    """Build headings + plain paragraphs with python-docx (no footnote yet)."""
    document = Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("This is the first paragraph of the document.")
    document.add_paragraph("This paragraph has a footnote reference.")
    document.add_heading("Background", level=2)
    document.add_paragraph("")
    document.add_paragraph("Final paragraph.")
    document.save(path)


def _footnotes_xml_bytes(footnotes_content: Sequence[tuple[str, str]]) -> bytes:
    """Build `word/footnotes.xml`: Word's two boilerplate separator footnotes
    (never referenced from the body - see
    [specs/get_footnotes.md §5](../../src/docx_mcp/specs/get_footnotes.md#5-limitations-non-goals))
    plus one real `w:footnote` per `(id, text)` pair in `footnotes_content`, in
    the order given."""
    footnotes = etree.Element(_w("footnotes"), nsmap={"w": WORD_NS})

    def _boilerplate_footnote(footnote_id: str, kind: str) -> None:
        footnote = etree.SubElement(footnotes, _w("footnote"))
        footnote.set(_w("type"), kind)
        footnote.set(_w("id"), footnote_id)
        paragraph = etree.SubElement(footnote, _w("p"))
        run = etree.SubElement(paragraph, _w("r"))
        run_props = etree.SubElement(run, _w("rPr"))
        etree.SubElement(run_props, _w("vertAlign")).set(_w("val"), "superscript")
        if kind == "separator":
            etree.SubElement(run, _w("separator"))
        else:
            etree.SubElement(run, _w("continuationSeparator"))

    _boilerplate_footnote("-1", "separator")
    _boilerplate_footnote("0", "continuationSeparator")

    for footnote_id, text in footnotes_content:
        footnote = etree.SubElement(footnotes, _w("footnote"))
        footnote.set(_w("id"), footnote_id)
        paragraph = etree.SubElement(footnote, _w("p"))
        ref_run = etree.SubElement(paragraph, _w("r"))
        ref_run_props = etree.SubElement(ref_run, _w("rPr"))
        etree.SubElement(ref_run_props, _w("rStyle")).set(_w("val"), "FootnoteReference")
        etree.SubElement(ref_run, _w("footnoteRef"))
        text_run = etree.SubElement(paragraph, _w("r"))
        text_elem = etree.SubElement(text_run, _w("t"))
        text_elem.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        text_elem.text = f" {text}"

    return etree.tostring(footnotes, xml_declaration=True, encoding="UTF-8", standalone=True)


def _add_footnote_references_to_document_xml(
    document_xml: bytes, anchors: Sequence[tuple[str, str]]
) -> bytes:
    """Append one `w:footnoteReference` run per `(paragraph_text, footnote_id)`
    in `anchors`, to the paragraph whose text matches `paragraph_text` exactly.
    Multiple pairs naming the same paragraph text append multiple references
    to it, in list order - this is how `footnotes.docx` places two footnote
    references in one paragraph (see
    [Roadmap.md, Phase 3](../../Roadmap.md#phase-3--footnotes-read-path-get_footnotes))."""
    root = etree.fromstring(document_xml)
    nsmap = {"w": WORD_NS}
    paragraphs = root.findall(".//w:body/w:p", namespaces=nsmap)

    def _paragraph_text(paragraph: etree._Element) -> str:
        return "".join(t.text or "" for t in paragraph.findall(".//w:t", namespaces=nsmap))

    for paragraph_text, footnote_id in anchors:
        target = next(p for p in paragraphs if _paragraph_text(p) == paragraph_text)
        run = etree.SubElement(target, _w("r"))
        run_props = etree.SubElement(run, _w("rPr"))
        etree.SubElement(run_props, _w("rStyle")).set(_w("val"), "FootnoteReference")
        ref = etree.SubElement(run, _w("footnoteReference"))
        ref.set(_w("id"), footnote_id)

    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _add_footnotes_content_type(content_types_xml: bytes) -> bytes:
    root = etree.fromstring(content_types_xml)
    override = etree.SubElement(root, f"{{{CONTENT_TYPES_NS}}}Override")
    override.set("PartName", "/word/footnotes.xml")
    override.set("ContentType", FOOTNOTES_CONTENT_TYPE)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _add_footnotes_relationship(rels_xml: bytes) -> bytes:
    root = etree.fromstring(rels_xml)
    existing_ids = {rel.get("Id") for rel in root}
    new_id = next(f"rId{n}" for n in range(1000, 2000) if f"rId{n}" not in existing_ids)
    relationship = etree.SubElement(root, f"{{{RELS_NS}}}Relationship")
    relationship.set("Id", new_id)
    relationship.set("Type", FOOTNOTES_RELATIONSHIP_TYPE)
    relationship.set("Target", "footnotes.xml")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _splice_in_footnotes(
    base_path: Path,
    output_path: Path,
    *,
    anchors: Sequence[tuple[str, str]],
    footnotes_content: Sequence[tuple[str, str]],
) -> None:
    """Rewrite `base_path` into `output_path`, adding a real `word/footnotes.xml`
    part, its `[Content_Types].xml` override and relationship, and one
    `w:footnoteReference` run per `anchors` entry. `anchors` is
    `(paragraph_text, footnote_id)` pairs consumed by
    `_add_footnote_references_to_document_xml`; `footnotes_content` is
    `(footnote_id, text)` pairs consumed by `_footnotes_xml_bytes`."""
    with zipfile.ZipFile(base_path) as source:
        parts = {name: source.read(name) for name in source.namelist()}

    parts["word/document.xml"] = _add_footnote_references_to_document_xml(
        parts["word/document.xml"], anchors
    )
    parts["[Content_Types].xml"] = _add_footnotes_content_type(parts["[Content_Types].xml"])
    parts["word/_rels/document.xml.rels"] = _add_footnotes_relationship(
        parts["word/_rels/document.xml.rels"]
    )
    parts["word/footnotes.xml"] = _footnotes_xml_bytes(footnotes_content)

    with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def generate_minimal_docx() -> Path:
    """Build `tests/fixtures/minimal.docx`: two heading levels, a footnote, plain paragraphs.

    Expected `read_document` output (see `tests/test_document.py`):

        # Introduction
        This is the first paragraph of the document.
        This paragraph has a footnote reference.[^1]
        ## Background

        Final paragraph.
    """
    output_path = FIXTURES_DIR / "minimal.docx"
    base_path = FIXTURES_DIR / "_minimal_base.docx"
    try:
        _build_base_docx(base_path)
        _splice_in_footnotes(
            base_path,
            output_path,
            anchors=[("This paragraph has a footnote reference.", "1")],
            footnotes_content=[("1", "This is a footnote.")],
        )
    finally:
        base_path.unlink(missing_ok=True)
    return output_path


def generate_structured_docx() -> Path:
    """Build `tests/fixtures/structured.docx` for `get_structure`/`get_metadata` (Phase 2).

    Three heading levels, including one *custom-named* style (`MySectionHeading`)
    that is `w:basedOn` built-in `Heading 2` and declares no `w:outlineLvl` of its
    own - it must resolve to heading level 2 purely through the `w:basedOn` chain,
    the case [specs/get_structure.md §3](../../src/docx_mcp/specs/get_structure.md#3-output-schema)
    describes. Also a table with one horizontally merged cell (`w:gridSpan`), and
    deterministic core properties for `get_metadata`.

    Expected `get_structure` output (see `tests/test_structure.py`):

        paragraphs[0] = "Introduction"                        style_id=Heading1  level=1
        paragraphs[1] = "This is a plain paragraph under the introduction." style_id=None level=None
        paragraphs[2] = "Background"                          style_id=Heading2  level=2
        paragraphs[3] = "Custom styled heading text"    style_id=MySectionHeading level=2 (basedOn)
        paragraphs[4] = "Deep Dive"                            style_id=Heading3  level=3
        paragraphs[5] = "Final plain paragraph."               style_id=None level=None
        tables[0].rows = [["Header A", "Header B"], ["Value 1", "Value 2"], ["Merged Row"]]

    Expected `get_metadata` output:

        title="Structured Test Document", author="Test Author",
        created="2024-01-01T00:00:00Z", modified="2024-06-15T12:30:00Z",
        word_count=29 (19 from paragraphs + 10 from table cells)
    """
    document = Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("This is a plain paragraph under the introduction.")
    document.add_heading("Background", level=2)

    custom_style = document.styles.add_style("MySectionHeading", WD_STYLE_TYPE.PARAGRAPH)
    custom_style.base_style = document.styles["Heading 2"]
    document.add_paragraph("Custom styled heading text", style=custom_style)

    document.add_heading("Deep Dive", level=3)
    document.add_paragraph("Final plain paragraph.")

    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Value 1"
    table.cell(1, 1).text = "Value 2"
    merged_cell = table.cell(2, 0)
    merged_cell.merge(table.cell(2, 1))
    merged_cell.text = "Merged Row"

    document.core_properties.title = "Structured Test Document"
    document.core_properties.author = "Test Author"
    document.core_properties.created = datetime.datetime(2024, 1, 1, 0, 0, 0, tzinfo=datetime.UTC)
    document.core_properties.modified = datetime.datetime(
        2024, 6, 15, 12, 30, 0, tzinfo=datetime.UTC
    )

    output_path = FIXTURES_DIR / "structured.docx"
    document.save(output_path)
    return output_path


def _build_footnotes_base_docx(path: Path) -> None:
    """Build headings + plain paragraphs for `footnotes.docx` (no footnotes yet)."""
    document = Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("This is a plain paragraph.")
    document.add_paragraph("This paragraph has two footnotes.")
    document.add_heading("Background", level=2)
    document.add_paragraph("This paragraph has a third footnote.")
    document.add_paragraph("Final paragraph.")
    document.save(path)


def generate_footnotes_docx() -> Path:
    """Build `tests/fixtures/footnotes.docx` for `get_footnotes` (Phase 3).

    Three real footnotes: ids `1` and `2` both anchored in paragraph 2 - the
    "two footnotes in the same paragraph" edge case
    [Roadmap.md, Phase 3](../../Roadmap.md#phase-3--footnotes-read-path-get_footnotes)
    calls out as most likely to break a naive anchor-to-paragraph mapping -
    and id `3` anchored in a separate paragraph, 4.

    Expected `get_structure`/`read_document` paragraph layout (see
    `tests/test_footnotes.py`, `tests/test_get_footnotes_tool.py`):

        paragraphs[0] = "Introduction"                          (Heading1)
        paragraphs[1] = "This is a plain paragraph."
        paragraphs[2] = "This paragraph has two footnotes."      <- anchors id=1, id=2
        paragraphs[3] = "Background"                             (Heading2)
        paragraphs[4] = "This paragraph has a third footnote."   <- anchor id=3
        paragraphs[5] = "Final paragraph."

    Expected `get_footnotes` output (content carries the leading space Word
    itself writes after the auto-number in the footnote's text run - the same
    convention `minimal.docx`'s existing single footnote already uses,
    preserved verbatim per `xml:space="preserve"` rather than trimmed):

        footnotes = [
            {"id": "1", "paragraph_index": 2, "content": " First footnote."},
            {"id": "2", "paragraph_index": 2, "content": " Second footnote."},
            {"id": "3", "paragraph_index": 4, "content": " Third footnote."},
        ]
    """
    output_path = FIXTURES_DIR / "footnotes.docx"
    base_path = FIXTURES_DIR / "_footnotes_base.docx"
    try:
        _build_footnotes_base_docx(base_path)
        _splice_in_footnotes(
            base_path,
            output_path,
            anchors=[
                ("This paragraph has two footnotes.", "1"),
                ("This paragraph has two footnotes.", "2"),
                ("This paragraph has a third footnote.", "3"),
            ],
            footnotes_content=[
                ("1", "First footnote."),
                ("2", "Second footnote."),
                ("3", "Third footnote."),
            ],
        )
    finally:
        base_path.unlink(missing_ok=True)
    return output_path


if __name__ == "__main__":
    shutil.rmtree(FIXTURES_DIR / "__pycache__", ignore_errors=True)
    minimal_path = generate_minimal_docx()
    print(f"wrote {minimal_path}")
    structured_path = generate_structured_docx()
    print(f"wrote {structured_path}")
    footnotes_path = generate_footnotes_docx()
    print(f"wrote {footnotes_path}")
